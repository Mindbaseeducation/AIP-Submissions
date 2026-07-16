import streamlit as st
import pandas as pd
import urllib.parse
from openpyxl import load_workbook, Workbook
from io import BytesIO
import zipfile
from copy import copy

# -------------------------------
# PAGE CONFIG
# -------------------------------
st.set_page_config(page_title="AIP Link Generator + Mentor Splitter", layout="wide")
st.title("📚 AIP Prefill Link Generator + Mentor-wise Splitter")
st.caption(
    "Upload the raw student file → prefilled MS Form links are generated → "
    "the file is split into mentor-wise workbooks (formatting + hyperlinks preserved)."
)

# -------------------------------
# CONFIG
# -------------------------------
BASE_URL = (
    "https://forms.office.com/Pages/ResponsePage.aspx?"
    "id=V15QOi70rU6qMklaCIJ0LfgKMwrSp1hDnVNgoqEsBV5UNjRVNEJaWUhLNUhRRkZHMDJYTVpXVzFPTiQlQCN0PWcu"
)

LINK_COL = "Prefilled Form Link"

# Columns required for link encoding
ENCODE_COLS = [
    "StudentID",
    "Student Name",
    "First strategic intervention",
    "Second strategic intervention",
    "Third strategic intervention",
]

# Columns required for mentor splitting
SPLIT_COLS = ["Mentor Name", "Regional Manager Name", "StudentID"]


# ===============================================================
# STEP 1 — LINK ENCODER
# ===============================================================
def encode(value):
    if pd.isna(value):
        return ""
    return urllib.parse.quote(str(value).strip())


def generate_url(row):
    return (
        BASE_URL
        + "&r187e87c4ce9f4cc3b8d497322e9887ef=" + encode(row["StudentID"])
        + "&r93fb97e6c3554faaa12db699650494a0=" + encode(row["Student Name"])
        + "&r76605a8e88574c03ad9d794239041a03=" + encode(row["First strategic intervention"])
        + "&re5ba2eea4f3140bca728577c76498278=" + encode(row["Second strategic intervention"])
        + "&rbf78e564c22243aca99debf0c0bbb296=" + encode(row["Third strategic intervention"])
    )


def add_links_to_workbook(original_wb, sheet_name, df):
    """
    Adds the prefilled-link column to the ORIGINAL workbook *in place* so that all
    of the source file's formatting is kept. The link cell shows the StudentID as
    display text and points to the generated MS Form URL (same behaviour as the
    standalone link encoder).

    Returns a df copy with the same column appended at the end, so the DataFrame
    column order stays aligned with the worksheet column order (the splitter maps
    columns by position).
    """
    ws = original_wb[sheet_name]

    # StudentID -> generated URL
    df = df.copy()
    url_map = {str(r["StudentID"]).strip(): generate_url(r) for _, r in df.iterrows()}

    # Locate the StudentID column in the sheet
    id_col_idx = None
    for c in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=c).value == "StudentID":
            id_col_idx = c
            break
    if id_col_idx is None:
        raise ValueError("StudentID column not found in the sheet.")

    # New column appended at the end of the sheet
    link_col_idx = ws.max_column + 1

    # Header cell (copy the StudentID header's look for consistency)
    hdr_src = ws.cell(row=1, column=id_col_idx)
    hdr_tgt = ws.cell(row=1, column=link_col_idx, value=LINK_COL)
    try:
        hdr_tgt.font = copy(hdr_src.font)
        hdr_tgt.fill = copy(hdr_src.fill)
        hdr_tgt.border = copy(hdr_src.border)
        hdr_tgt.alignment = copy(hdr_src.alignment)
        hdr_tgt.number_format = hdr_src.number_format
    except Exception:
        pass

    # Data rows: display StudentID, link to the form
    for r in range(2, ws.max_row + 1):
        sid_val = ws.cell(row=r, column=id_col_idx).value
        if sid_val is None:
            continue
        sid = str(sid_val).strip()
        url = url_map.get(sid)
        cell = ws.cell(row=r, column=link_col_idx, value=sid)
        if url:
            cell.hyperlink = url
            cell.style = "Hyperlink"

    # Keep the DataFrame aligned with the sheet (link col value = StudentID text)
    df[LINK_COL] = df["StudentID"].apply(lambda x: str(x).strip())
    return df


# ===============================================================
# STEP 2 — MENTOR-WISE FORMATTED SPLIT
# ===============================================================
def df_to_formatted_workbook(original_wb, original_ws_name, df_filtered, id_col_name):
    original_ws = original_wb[original_ws_name]

    # Find ID column index
    id_col_idx = None
    for c in range(1, original_ws.max_column + 1):
        if original_ws.cell(row=1, column=c).value == id_col_name:
            id_col_idx = c
            break

    if id_col_idx is None:
        raise ValueError(f"ID column '{id_col_name}' not found.")

    # Map ID → original row
    id_to_row = {
        str(original_ws.cell(row=r, column=id_col_idx).value): r
        for r in range(2, original_ws.max_row + 1)
        if original_ws.cell(row=r, column=id_col_idx).value is not None
    }

    # Create new workbook
    new_wb = Workbook()
    new_ws = new_wb.active
    new_ws.title = "Students"

    # Copy column widths
    for col, dim in original_ws.column_dimensions.items():
        if dim.width:
            new_ws.column_dimensions[col].width = dim.width

    columns = list(df_filtered.columns)

    # Copy headers with formatting
    for col_idx, col_name in enumerate(columns, start=1):
        src = original_ws.cell(row=1, column=col_idx)
        tgt = new_ws.cell(row=1, column=col_idx, value=col_name)

        try:
            tgt.font = copy(src.font)
            tgt.fill = copy(src.fill)
            tgt.border = copy(src.border)
            tgt.alignment = copy(src.alignment)
            tgt.number_format = src.number_format
        except Exception:
            pass

    # Fill rows
    for row_idx, (_, row) in enumerate(df_filtered.iterrows(), start=2):

        orig_row = id_to_row.get(str(row.get(id_col_name)), 2)

        # Copy row height
        if original_ws.row_dimensions.get(orig_row):
            new_ws.row_dimensions[row_idx].height = original_ws.row_dimensions[orig_row].height

        for col_idx, col_name in enumerate(columns, start=1):
            val = row[col_name]
            src = original_ws.cell(row=orig_row, column=col_idx)
            tgt = new_ws.cell(row=row_idx, column=col_idx, value=val)

            # Copy styles
            try:
                tgt.font = copy(src.font)
                tgt.fill = copy(src.fill)
                tgt.border = copy(src.border)
                tgt.alignment = copy(src.alignment)
                tgt.number_format = src.number_format
            except Exception:
                pass

            # 🔥 Preserve hyperlink EXACTLY (no modification)
            if src.hyperlink:
                tgt.hyperlink = src.hyperlink.target
                tgt.style = "Hyperlink"

    return new_wb


# ===============================================================
# MAIN APP
# ===============================================================
uploaded_file = st.file_uploader("Upload Excel file", type=["xlsx"])

if uploaded_file:
    # Read the bytes once, reuse for both pandas and openpyxl
    file_bytes = uploaded_file.getvalue()

    df = pd.read_excel(BytesIO(file_bytes))

    # Validate required columns (both stages)
    missing = sorted({c for c in (ENCODE_COLS + SPLIT_COLS) if c not in df.columns})
    if missing:
        st.error("❌ Missing required column(s): " + ", ".join(missing))
        st.stop()

    st.success("✅ File uploaded successfully!")

    # ---- STEP 1: generate prefilled links into the original workbook ----
    original_wb = load_workbook(BytesIO(file_bytes))
    sheet_name = original_wb.sheetnames[0]
    df = add_links_to_workbook(original_wb, sheet_name, df)

    with st.expander("🔍 Preview generated links"):
        preview = df.copy()
        preview["Generated URL"] = preview.apply(generate_url, axis=1)
        st.dataframe(
            preview[["StudentID", "Student Name", "Generated URL"]].head(20),
            use_container_width=True,
        )

    # Filename label (defaults to July, editable for other months)
    label = st.text_input("Filename label (e.g. month)", value="July")

    # ---- STEP 2: filters ----
    regional_managers = ["All"] + sorted(df["Regional Manager Name"].dropna().unique())
    selected_regional_manager = st.selectbox("Select Regional Manager Name", regional_managers)

    filtered_df = (
        df
        if selected_regional_manager == "All"
        else df[df["Regional Manager Name"] == selected_regional_manager]
    )

    mentors = sorted(filtered_df["Mentor Name"].dropna().unique())
    selected_mentor = st.selectbox("Select Mentor", mentors)

    final_df = filtered_df[filtered_df["Mentor Name"] == selected_mentor]

    st.write(f"### Students under **{selected_mentor}** ({len(final_df)} students)")
    st.dataframe(final_df, use_container_width=True)

    # ---- SINGLE MENTOR DOWNLOAD ----
    try:
        wb = df_to_formatted_workbook(
            original_wb,
            sheet_name,
            final_df,
            id_col_name="StudentID",
        )

        buffer = BytesIO()
        wb.save(buffer)

        st.download_button(
            label=f"📥 Download Excel for {selected_mentor}",
            data=buffer.getvalue(),
            file_name=f"{selected_mentor} AIP Follow-up Student List ({label}).xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    except Exception as e:
        st.error(f"❌ Error: {e}")

    # ---- ZIP DOWNLOAD (ALL MENTORS) ----
    st.write("---")
    st.write("### 📦 Download ZIP (All Mentors)")

    zip_buffer = BytesIO()

    with zipfile.ZipFile(zip_buffer, "w") as zf:
        for mentor in mentors:
            try:
                m_df = filtered_df[filtered_df["Mentor Name"] == mentor]

                wb = df_to_formatted_workbook(
                    original_wb,
                    sheet_name,
                    m_df,
                    id_col_name="StudentID",
                )

                mb = BytesIO()
                wb.save(mb)

                zf.writestr(
                    f"{mentor} AIP Follow-up Student List ({label}).xlsx",
                    mb.getvalue(),
                )

            except Exception as e:
                zf.writestr(f"{mentor}_ERROR.txt", str(e))

    st.download_button(
        label="📥 Download ZIP",
        data=zip_buffer.getvalue(),
        file_name="Mentor_Files.zip",
        mime="application/zip",
    )
